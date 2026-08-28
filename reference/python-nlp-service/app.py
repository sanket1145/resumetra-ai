"""FastAPI NLP service: text extraction, resume/job parsing and match scoring."""

from __future__ import annotations

import io

import pdfplumber
from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from pydantic import BaseModel

from nlp_core import match, parse_job, parse_resume

app = FastAPI(title="Resume NLP Service", version="1.0.0")


class TextPayload(BaseModel):
    text: str


class MatchPayload(BaseModel):
    resume_text: str
    job_text: str


def pdf_to_text(data: bytes) -> str:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def docx_to_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


@app.get("/health")
def health() -> dict:
    return {"status": "ok"}


@app.post("/extract")
async def extract(file: UploadFile = File(...)) -> dict:
    data = await file.read()
    name = (file.filename or "").lower()
    if name.endswith(".pdf"):
        text = pdf_to_text(data)
    elif name.endswith(".docx"):
        text = docx_to_text(data)
    else:
        raise HTTPException(status_code=400, detail="Only PDF and DOCX files are supported.")
    if not text.strip():
        raise HTTPException(status_code=422, detail="No selectable text found in this document.")
    return {"text": text, "characters": len(text)}


@app.post("/parse-resume")
def parse_resume_endpoint(payload: TextPayload) -> dict:
    return parse_resume(payload.text)


@app.post("/parse-job")
def parse_job_endpoint(payload: TextPayload) -> dict:
    return parse_job(payload.text)


@app.post("/match")
def match_endpoint(payload: MatchPayload) -> dict:
    resume = parse_resume(payload.resume_text)
    job = parse_job(payload.job_text)
    return {"resume": resume, "job": job, "analysis": match(resume, job)}
